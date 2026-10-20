#
# This module holds the Excel report module
#
# Design decisions:
# - Needs the optional `openpyxl` package (`pip install fixturerig[excel]`).
#   It is imported when the module is created, so a missing package is
#   reported as a configuration error before any test runs.
# - Sheets:
#     * "Summary": one row per test (group path, name, verdict, message,
#       duration, class, traceability), with an autofilter for triage.
#     * "Details": everything tests reported, in order. Simple tables are
#       written as real cell ranges; complex tables are flattened to text.
#     * "Charts": the data of every reported chart, with a native Excel
#       chart next to it.
#

from __future__ import annotations

from pathlib import Path

from fixturerig.base.verdict import VerdictStatus
from fixturerig.report.elements import Chart, Image, Table, Text
from fixturerig.report.events import VerdictEvent
from fixturerig.report.recording import RecordingReport, RunRecord

_FILLS = {
    VerdictStatus.PASSED: "C6EFCE",
    VerdictStatus.FAILED: "FFC7CE",
    VerdictStatus.ERROR: "FF9A9A",
    VerdictStatus.SKIPPED: "FFEB9C",
    VerdictStatus.NONE: "E4D3EE",
    VerdictStatus.KNOWNFAILED: "CFE8EC",
}


def _require_openpyxl():
    try:
        import openpyxl
    except ImportError as exc:
        raise ImportError(
            "ExcelReport needs the 'openpyxl' package: pip install fixturerig[excel]"
        ) from exc
    return openpyxl


class ExcelReport(RecordingReport):
    """Writes the results and reported content to an .xlsx workbook.

    :param path: the workbook to write.
    """

    def __init__(self, path: Path = Path("fixturerig-report.xlsx"), **kwargs):
        super().__init__(**kwargs)
        _require_openpyxl()
        self.path = Path(path)

    def write(self, run: RunRecord) -> None:
        openpyxl = _require_openpyxl()
        from openpyxl.styles import Font, PatternFill

        workbook = openpyxl.Workbook()
        bold = Font(bold=True)

        summary = workbook.active
        summary.title = "Summary"
        summary.append(["Group", "Test", "Verdict", "Message", "Duration [s]", "Class", "Traceability"])
        for cell in summary[1]:
            cell.font = bold
        for test in run.tests():
            spec, done = test.started, test.finished
            status = done.verdict.status if done else VerdictStatus.NONE
            summary.append(
                [
                    " / ".join(spec.path[:-1]),
                    spec.name,
                    status.value,
                    done.verdict.message if done else "test did not finish",
                    round(done.duration, 3) if done else None,
                    spec.cls,
                    ", ".join(spec.traceability),
                ]
            )
            summary.cell(row=summary.max_row, column=3).fill = PatternFill("solid", fgColor=_FILLS[status])
        summary.auto_filter.ref = summary.dimensions
        summary.freeze_panes = "A2"
        for column, width in zip("ABCDEFG", (24, 28, 12, 50, 12, 36, 24)):
            summary.column_dimensions[column].width = width

        details = workbook.create_sheet("Details")
        details.append(["Test", "Kind", "Content"])
        for cell in details[1]:
            cell.font = bold
        details.column_dimensions["A"].width = 28
        details.column_dimensions["B"].width = 10
        details.column_dimensions["C"].width = 80

        charts_sheet = None
        chart_row = 1

        events = [event for test in run.tests() for event in test.content]
        for event in events:
            test_name = " / ".join(event.path)
            element = event.element
            if isinstance(event, VerdictEvent):
                details.append([test_name, "verdict", f"{event.verdict.status.value}: {event.verdict.message}"])
                details.cell(row=details.max_row, column=3).fill = PatternFill(
                    "solid", fgColor=_FILLS[event.verdict.status]
                )
            elif isinstance(element, Table) and element.is_simple:
                details.append([test_name, "table", element.title or ""])
                if element.headers:
                    details.append([None, None, *element.headers])
                    for cell in details[details.max_row][2:]:
                        cell.font = bold
                for row in element.rows:
                    details.append([None, None, *row])
            elif isinstance(element, Chart):
                details.append([test_name, "chart", element.to_text()])
                if charts_sheet is None:
                    charts_sheet = workbook.create_sheet("Charts")
                chart_row = _write_chart(charts_sheet, element, chart_row, bold) + 2
            elif isinstance(element, (Text, Table, Image)):
                details.append([test_name, event.kind, element.to_text()])

        self.path.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(self.path)


def _write_chart(sheet, chart: Chart, row: int, bold) -> int:
    """Write the chart data at `row` and add an Excel chart; returns the next free row."""
    from openpyxl.chart import BarChart, LineChart, Reference, ScatterChart
    from openpyxl.chart import Series as XlSeries

    sheet.cell(row=row, column=1, value=chart.title).font = bold
    first_data_row = row + 1
    column = 1
    longest = 0
    references = []
    for series in chart.series:
        sheet.cell(row=first_data_row, column=column, value=chart.x_label or "x").font = bold
        sheet.cell(row=first_data_row, column=column + 1, value=series.name).font = bold
        for offset, (x, y) in enumerate(zip(series.x, series.y), start=1):
            sheet.cell(row=first_data_row + offset, column=column, value=x)
            sheet.cell(row=first_data_row + offset, column=column + 1, value=y)
        references.append((column, len(series.y), series.name))
        longest = max(longest, len(series.y))
        column += 2

    if chart.kind == "bar":
        xl_chart = BarChart()
    elif chart.kind == "line":
        xl_chart = LineChart() if len(chart.series) <= 1 else ScatterChart()
    else:
        xl_chart = ScatterChart()
    xl_chart.title = chart.title
    xl_chart.x_axis.title = chart.x_label
    xl_chart.y_axis.title = chart.y_label

    for col, count, name in references:
        if count == 0:
            continue
        x_ref = Reference(sheet, min_col=col, min_row=first_data_row + 1, max_row=first_data_row + count)
        y_ref = Reference(sheet, min_col=col + 1, min_row=first_data_row, max_row=first_data_row + count)
        if isinstance(xl_chart, ScatterChart):
            xl_series = XlSeries(y_ref, x_ref, title_from_data=True)
            if chart.kind == "scatter":
                xl_series.marker.symbol = "circle"
                xl_series.graphicalProperties.line.noFill = True
            xl_chart.series.append(xl_series)
        else:
            xl_chart.add_data(y_ref, titles_from_data=True)
            xl_chart.set_categories(x_ref)

    if references and longest:
        sheet.add_chart(xl_chart, sheet.cell(row=row, column=column + 1).coordinate)
    # A chart anchored next to the data occupies roughly 15 rows.
    return first_data_row + max(longest, 15)
